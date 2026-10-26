"""הסטים לבדיקות של יום התשלום (docs/before-paid-tests.md סעיף 5): 30 שאלות מכל סט, ו-10 הודעות עם כמה שאלות.

- data/questions_dev30.csv: 30 מהסט הראשי (questions.csv) - התשובות הצפויות שלו כבר אושרו. לפיתוח ול-A/B.
- data/questions_dev30b.csv: 30 אחרות מהסט הראשי, בלי אף שאלה מ-dev30 - הסבב השני של ה-A/B (החלטת הבעלים 26/10: הריצה
  החוזרת על שאלות אחרות). שינוי שעוזר בשני הסבבים - לא מקרה של 30 שאלות מסוימות.
- data/questions_final30.csv: 30 מהסט החדש (questions_v2.csv) - הבוט לא מכיר אותן. לבדיקה הסופית; צריך אישור.
- data/questions_multi.csv: 10 הודעות עם 2-3 שאלות (MULTI_ANSWER), בנויות משאלות של הסט הראשי: התשובה הצפויה היא
  התשובות שכבר אושרו, מילה במילה, חלק אחרי חלק. צריך לאשר רק את ניסוח ההודעה.

הבחירה מכל סוג - לפי היחס בסט (84 -> 30: 15 עם תשובה, 4 אישיות, 4 לא באתר, 3 עקיפה, 2 הבהרה, 1 ברכה, 1 לא קשורה),
ובתוך הסוג - הגרלה עם זרע קבוע (SEED): אותה בחירה בכל הרצה, ובלי לבחור שאלות שהבוט "טוב בהן". שאלות זהות - פעם אחת.
השורות מועתקות כמו שהן (בלי שינוי בתשובות הצפויות), ועם orig_n - המספר שלהן בקובץ המקור.

קובץ בדיקה (Excel) לסט הסופי ולשאלות עם כמה חלקים: לכל שאלה עם עמוד מקור - המשפטים מהעמוד שהכי קרובים לתשובה הצפויה,
ובדיקה שכל מספר בתשובה הצפויה כתוב בעמוד. עמודות "אישור" ו"הערות" - למילוי.

הרצה:  python build_question_sets.py [--review "Claude outputs/review_questions.xlsx"]
"""
import argparse
import csv
import json
import random
import re
import sys
from pathlib import Path
from urllib.parse import unquote

from phoenix import answer as A
from phoenix import bm25, config
from phoenix.chunk import CHUNKS_FILE

DATA = config.ROOT / "data"
SEED = 26102026
QUOTAS = {"answerable": 15, "personal": 4, "not_on_site": 4, "injection": 3, "clarify": 2, "greeting": 1,
          "off_topic": 1}
FIELDS = ["question", "expected_answer", "source_url", "type", "orig_n"]
# הודעה עם כמה שאלות: (הניסוח, מספרי השאלות ב-questions.csv). התשובה הצפויה - שלהן, לפי הסדר. 68 - כבר בסט, כמו שהיא.
MULTI = [
    ("הרכב שלי הוכרז טוטאל לוס. תוך כמה זמן מקבלים את הכסף, ולכמה ימים מקבלים רכב חליפי?", [8, 9]),
    ("רכב שמבוטח בהפניקס פגע ברכב שלי. מה עושים, וכמה זמן יש לי לתבוע?", [11, 12]),
    ("איך מגישים תביעה בביטוח בריאות, ועד מתי אפשר להגיש אותה?", [16, 15]),
    ("יש לי הצפה בבית. מה עושים קודם כל, ועד מתי אפשר להגיש תביעה בביטוח דירה?", [3, 6]),
    (None, [68]),
    ("איך אני מבטלת את ביטוח הרכב שלי, ואיך מגישים תביעה על ביטוח הדירה?", [75, 73]),
    ("אני שוכרת דירה. אפשר לעשות אצלכם ביטוח תכולה, ויש הנחה בביטוח דירה לגמלאים?", [1, 29]),
    ("מה זה השתתפות עצמית, והאם הביטוח שלי כולל רכב חלופי?", [69, 45]),
    ("איך מגישים תביעה בביטוח רכב, איך מבטלים את ביטוח הרכב, ומה זה השתתפות עצמית?", [60, 75, 69]),
    ("מה צריך כדי למשוך כסף מקרן השתלמות, ומה הגיל המינימלי לקבלת קצבת זקנה בקרן הפנסיה המקיפה של הפניקס?",
     [81, 83]),
]
MULTI_NOTES = {5: "השאלה מהסט הראשי, כמו שהיא", 6: "שני תחומים", 7: "חלק 2 לא באתר", 8: "חלק 2 אישי",
               9: "שלושה חלקים", 1: "שני חלקים מאותו עמוד", 2: "שני חלקים מאותו עמוד"}


def read(path: Path) -> list[dict]:
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def stratified(rows: list[dict], quotas: dict[str, int], seed: int = SEED,
               exclude: set[str] = frozenset()) -> list[dict]:
    """מכל סוג - quotas[סוג] שאלות בהגרלה עם זרע קבוע, בלי כפולות (אותו נוסח), לפי הסדר בקובץ המקור. exclude - נוסחים
    שכבר בסט אחר (אחרי " ".join(split()))."""
    rng = random.Random(seed)
    chosen = []
    for kind, count in quotas.items():
        seen, pool = set(exclude), []
        for n, row in enumerate(rows, 1):
            text = " ".join(row["question"].split())
            if row["type"] == kind and text not in seen:
                seen.add(text)
                pool.append(row | {"orig_n": n})
        if len(pool) < count:
            raise SystemExit(f"{kind}: רק {len(pool)} שאלות, ביקשנו {count}")
        chosen += rng.sample(pool, count)
    return sorted(chosen, key=lambda r: r["orig_n"])


def multi_rows(rows: list[dict]) -> list[dict]:
    """הודעות עם כמה שאלות: הניסוח מ-MULTI, והתשובה הצפויה - התשובות שאושרו לכל שאלה, מילה במילה, חלק אחרי חלק."""
    out = []
    for text, numbers in MULTI:
        parts = [rows[n - 1] for n in numbers]
        if text is None:  # שאלה שכבר בסט - כמו שהיא
            out.append(parts[0] | {"orig_n": str(numbers[0])})
            continue
        expected = "\n".join(f"({i}) {p['question']}\n{p['expected_answer']}" for i, p in enumerate(parts, 1))
        urls = " ".join(dict.fromkeys(u for p in parts for u in p["source_url"].split()))
        out.append({"question": text, "expected_answer": expected, "source_url": urls, "type": "answerable",
                    "orig_n": "+".join(map(str, numbers)),
                    # לבדיקה בקובץ Excel בלבד (לא נכתב ל-CSV): התשובות עצמן, בלי "(1)" ובלי השאלות - אחרת "1" ו-"2"
                    # של הסימון היו נבדקים כמספרים מהאתר
                    "check_text": "\n".join(p["expected_answer"] for p in parts),
                    "with_contact": any(p["type"] in ("personal", "not_on_site") for p in parts)})
    return out


def write(path: Path, rows: list[dict]):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


STEP_NUMBER = re.compile(r"(?<![\w.])\d{1,2}\)(?=\s)")  # "חובה: 1) ... 2) ..." - מספור של שלבים, לא מספר מהאתר


def _page_key(url: str) -> str:
    return unquote(url.split("#")[0]).rstrip("/").lower()


def site_pages() -> dict[str, list[str]]:
    """כתובת העמוד (בלי #page=N, בלי / בסוף) -> הטקסטים של הקטעים שלו."""
    pages: dict[str, list[str]] = {}
    for line in open(CHUNKS_FILE, encoding="utf-8"):
        chunk = json.loads(line)
        pages.setdefault(_page_key(chunk["url"]), []).append(chunk["text"])
    return pages


def _sentences(texts: list[str]) -> list[str]:
    return list(dict.fromkeys(s for t in texts
                              for s in A.passage_sentences(t.split("\n\n", 1)[-1], A.PDF_MARK in t[:200].lower())
                              if len(s.split()) >= 4))


def nearest_on_site(question: str, top: int = 2) -> str:
    """"לא באתר": מה שהכי קרוב לשאלה בכל האתר (BM25 - מילים משותפות), כדי לוודא שהוא לא עונה עליה."""
    found = bm25.search(question, n=3)
    words = set(bm25.tokenize(question))
    scored = sorted(((s, c["url"]) for c, _ in found for s in _sentences([c["text"]])),
                    key=lambda pair: -len(words & set(bm25.tokenize(pair[0]))))
    return "הכי קרוב באתר: " + " / ".join(f"{s} ({url})" for s, url in scored[:top]) if scored else ""


def support(row: dict, pages: dict[str, list[str]], site: str, top: int = 2) -> tuple[str, str]:
    """(המשפטים מהעמוד שהכי קרובים לתשובה הצפויה, בדיקת המספרים). קרוב = מספרים משותפים (משקל 3) ומילים משותפות.
    "לא באתר" - במקום עמוד המקור: מה שהכי קרוב לשאלה בכל האתר. במספרים של תשובה אישית או "לא באתר" - גם עמוד צור
    קשר (מספרי ההפניה כתובים שם)."""
    expected = row.get("check_text") or row["expected_answer"]
    numbers = A._numbers(STEP_NUMBER.sub("", expected))
    texts = [t for url in row["source_url"].split() for t in pages.get(_page_key(url), [])]
    contact = pages.get(_page_key(A.CONTACT_URL), []) if row.get("with_contact") or row["type"] in (
        "personal", "not_on_site") else []
    missing = A.unsupported_numbers(STEP_NUMBER.sub("", expected), texts + contact) if numbers else []
    nowhere = A.unsupported_numbers(" ".join(missing), [site]) if missing else []  # site - כל הטקסט של האתר
    elsewhere = [n for n in missing if n not in nowhere]
    check = "" if not numbers else "; ".join(filter(None, [
        "לא באתר בכלל: " + ", ".join(nowhere) if nowhere else "",
        "לא בעמוד, אבל באתר (בעמוד אחר): " + ", ".join(elsewhere) if elsewhere else ""])) or "כל המספרים בעמוד"
    if row["type"] == "not_on_site":
        return nearest_on_site(row["question"]), check
    if not texts:
        return ("", "העמוד לא באינדקס" if row["source_url"] else "")
    words = set(bm25.tokenize(expected))
    scored = sorted(_sentences(texts),
                    key=lambda s: -(3 * len(numbers & A._numbers(s)) + len(words & set(bm25.tokenize(s)))))
    return " / ".join(re.sub(r"^(?:#+|[-•*])\s+", "", s) for s in scored[:top]), check


def review(path: Path, final: list[dict], multi: list[dict]):
    """קובץ Excel לבדיקה: הסבר, הסט הסופי, ההודעות עם כמה חלקים. מימין לשמאל; "אישור" - רשימה נפתחת."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    pages = site_pages()
    site = "\n".join(t for texts in pages.values() for t in texts)
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill, fill_me = PatternFill("solid", fgColor="DDE7F3"), PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
    book = Workbook()
    legend = book.active
    legend.title = "הסבר"
    legend.sheet_view.rightToLeft = True
    lines = [
        ("מה בקובץ", "שני סטים לבדיקה ביום התשלום. בכל שורה - שאלה ותשובה צפויה, ולצידן המשפטים מהאתר שהכי קרובים "
                     "לתשובה הצפויה (נבחרו אוטומטית) ובדיקה אוטומטית שכל מספר בתשובה הצפויה כתוב בעמוד."),
        ("מה למלא", "רק העמודות הצהובות: \"אישור\" (כן / לא / תיקון) ו\"הערות\". אם צריך תיקון - לכתוב בהערות מה "
                    "נכון; אני לא משנה תשובה צפויה בעצמי."),
        ("סט סופי - 30", "30 שאלות מהסט החדש (questions_v2.csv), מכל הסוגים, בהגרלה עם זרע קבוע. הבוט לא מכיר אותן."),
        ("כמה חלקים - 10", "הודעות עם 2-3 שאלות, בנויות משאלות של הסט הראשי שכבר אישרת. התשובה הצפויה היא התשובות "
                           "שאישרת, מילה במילה - צריך לאשר רק את ניסוח ההודעה."),
        ("דוגמה למילוי", "אישור: תיקון | הערות: באתר כתוב 30 יום ולא 14 - לעדכן בתשובה הצפויה"),
    ]
    for r, (title, text) in enumerate(lines, 1):
        legend.cell(r, 1, title).font = bold
        legend.cell(r, 2, text).font = font
        legend.cell(r, 2).alignment = wrap
    legend.column_dimensions["A"].width, legend.column_dimensions["B"].width = 18, 110

    columns = [("#", 5), ("מספר במקור", 10), ("סוג", 11), ("שאלה", 38), ("תשובה צפויה", 60), ("עמוד המקור", 30),
               ("המשפטים מהאתר (אוטומטי)", 60), ("מספרים (אוטומטי)", 16), ("אישור", 9), ("הערות", 30)]
    for title, rows, note in (("סט סופי - 30", final, {}), ("כמה חלקים - 10", multi, MULTI_NOTES)):
        sheet = book.create_sheet(title)
        sheet.sheet_view.rightToLeft = True
        names = columns if not note else columns[:3] + [("מה נבדק", 16)] + columns[3:]
        for c, (name, width) in enumerate(names, 1):
            cell = sheet.cell(1, c, name)
            cell.font, cell.fill, cell.alignment = bold, head_fill, wrap
            sheet.column_dimensions[cell.column_letter].width = width
        approve = DataValidation(type="list", formula1='"כן,לא,תיקון"', allow_blank=True)
        sheet.add_data_validation(approve)
        for r, row in enumerate(rows, 2):
            sentences, check = support(row, pages, site)
            values = [r - 1, row["orig_n"], row["type"]] + ([note.get(r - 1, "")] if note else []) + [
                row["question"], row["expected_answer"], row["source_url"], sentences, check, "", ""]
            for c, value in enumerate(values, 1):
                cell = sheet.cell(r, c, value)
                cell.font, cell.alignment = font, wrap
            for c in (len(values) - 1, len(values)):
                sheet.cell(r, c).fill = fill_me
            approve.add(sheet.cell(r, len(values) - 1))
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--review", default=str(config.ROOT / "Claude outputs" / "review_questions.xlsx"))
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    main_rows, new_rows = read(DATA / "questions.csv"), read(DATA / "questions_v2.csv")
    dev, final, multi = stratified(main_rows, QUOTAS), stratified(new_rows, QUOTAS), multi_rows(main_rows)
    second = stratified(main_rows, QUOTAS, exclude={" ".join(r["question"].split()) for r in dev})
    for name, rows in (("questions_dev30.csv", dev), ("questions_dev30b.csv", second), ("questions_final30.csv", final),
                       ("questions_multi.csv", multi)):
        write(DATA / name, rows)
        print(f"data/{name}: {len(rows)} שאלות")
    review(Path(args.review), final, multi)
    print(f"קובץ הבדיקה: {args.review}")


if __name__ == "__main__":
    main()
